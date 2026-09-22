from pathlib import Path

import openpyxl

from logic import SIN_CODIGO_SAP, SIN_INCIDENCIA, procesar


def _crear_comercial(path: Path, filas: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name in ["Hoja 1", "Resumen", "Prov-Pedido"]:
        wb.create_sheet(sheet_name)
    ws = wb.create_sheet("Prov-Ref")
    ws.append(
        [
            "Código Prov", "Pedido", "Cod.Art.", "Ean Unidad Item", "Ean Caja",
            "Descripción", "Pedidas", "Servidas", "No Servidas", "Tasa", "Observaciones",
        ]
    )
    for fila in filas:
        codigo_prov, pedido, cod_art, ean, no_servidas = fila
        ws.append([codigo_prov, pedido, cod_art, ean, "", "desc", 1, 0, no_servidas, 0, None])
    wb.create_sheet("Hoja1")
    wb.save(path)


def _crear_sap(path: Path, filas: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(
        [
            "Referencia de cliente", "Fecha del documento", "Clase doc.ventas", "Documento de ventas",
            "Posición", "Solicitante", "Material", "Cantidad de pedido (Posición)", "Un.medida venta",
            "Valor neto (posición)", "Moneda del documento", "Descripción del motivo de rechazo",
            "Motivo de rechazo",
        ]
    )
    for pedido, material, motivo_desc, motivo_cod in filas:
        ws.append([pedido, None, "ZSTC", "1", "10", "1", material, 1, "UC", 1, "EUR", motivo_desc, motivo_cod])
    wb.save(path)


def _crear_cruce(path: Path, filas: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cruce_EAN_SAP"
    ws.append(["EAN", "Codigo_SAP", "Descripcion"])
    for ean, sap, desc in filas:
        ws.append([ean, sap, desc])
    wb.save(path)


def test_motivo_encontrado(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H100", "X1", "1111111111111", 5)])
    _crear_sap(tmp_path / "sap.xlsx", [("H100", "9999999", "Falta disponibilidad", "Z2")])
    _crear_cruce(tmp_path / "cruce.xlsx", [("1111111111111", "9999999", "producto")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.actualizadas == 1
    assert resultado.sin_incidencia == 0
    assert resultado.sin_codigo_sap == 0

    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == "Falta disponibilidad"


def test_ean_multiples_codigos_sap_uno_con_incidencia(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H200", "X1", "2222222222222", 3)])
    _crear_sap(
        tmp_path / "sap.xlsx",
        [("H200", "1000001", "", ""), ("H200", "1000002", "Cancelado por el cliente", "Z1")],
    )
    _crear_cruce(
        tmp_path / "cruce.xlsx",
        [("2222222222222", "1000001", "p"), ("2222222222222", "1000002", "p")],
    )

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.actualizadas == 1
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == "Cancelado por el cliente"


def test_sin_incidencia_registrada(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H300", "X1", "3333333333333", 2)])
    _crear_sap(tmp_path / "sap.xlsx", [("H300", "5000000", "", "")])
    _crear_cruce(tmp_path / "cruce.xlsx", [("3333333333333", "5000000", "p")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.sin_incidencia == 1
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == SIN_INCIDENCIA


def test_sin_incidencia_cuando_pedido_no_tiene_ese_material(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H400", "X1", "4444444444444", 1)])
    _crear_sap(tmp_path / "sap.xlsx", [("H400", "6000000", "otro motivo", "Z9")])
    _crear_cruce(tmp_path / "cruce.xlsx", [("4444444444444", "7000000", "p")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.sin_incidencia == 1
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == SIN_INCIDENCIA


def test_ean_no_esta_en_tabla_cruce(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H500", "X1", "9999999999999", 4)])
    _crear_sap(tmp_path / "sap.xlsx", [("H500", "1234567", "algun motivo", "Z1")])
    _crear_cruce(tmp_path / "cruce.xlsx", [("1111111111111", "1234567", "p")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.sin_codigo_sap == 1
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == SIN_CODIGO_SAP


def test_motivo_con_texto_personalizado(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H700", "X1", "5555555555555", 1)])
    _crear_sap(
        tmp_path / "sap.xlsx",
        [("H700", "8000000", "AC Falta  disponibilidad de la mercancía", "Z2")],
    )
    _crear_cruce(tmp_path / "cruce.xlsx", [("5555555555555", "8000000", "p")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.actualizadas == 1
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value == "Falta de disponibilidad informada"


def test_fila_no_servidas_cero_no_se_toca(tmp_path: Path):
    _crear_comercial(tmp_path / "comercial.xlsx", [("1", "H600", "X1", "1111111111111", 0)])
    _crear_sap(tmp_path / "sap.xlsx", [])
    _crear_cruce(tmp_path / "cruce.xlsx", [("1111111111111", "1234567", "p")])

    resultado = procesar(
        tmp_path / "comercial.xlsx", tmp_path / "sap.xlsx", tmp_path / "cruce.xlsx", tmp_path / "out.xlsx"
    )

    assert resultado.total_procesadas == 0
    wb = openpyxl.load_workbook(tmp_path / "out.xlsx")
    ws = wb["Prov-Ref"]
    assert ws.cell(row=2, column=11).value is None
