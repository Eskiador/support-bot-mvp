"""Registro Excel acumulado: una fila por PDF generado."""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

NOMBRE_ARCHIVO = "registro_confirmaciones.xlsx"
HOJA = "Registro"
COLUMNAS = [
    "Fecha proceso",
    "Archivo generado",
    "Nº documento",
    "Cliente",
    "GLN emisor",
    "Tipo",
    "Fecha documento",
    "Nº pedido",
    "Páginas",
    "Páginas en el combinado",
    "Combinado original",
    "Carpeta destino",
]


class ErrorRegistro(Exception):
    pass


@dataclass
class FilaPrevia:
    fecha_proceso: str
    archivo: str
    num_doc: str
    gln: str


class Registro:
    def __init__(self, carpeta: Path):
        self.carpeta = carpeta
        self.ruta = carpeta / NOMBRE_ARCHIVO

    def comprobar_escribible(self) -> None:
        """Falla si el Excel está abierto (Windows lo bloquea) o no se puede crear."""
        try:
            self.carpeta.mkdir(parents=True, exist_ok=True)
            if self.ruta.exists():
                with open(self.ruta, "r+b"):
                    pass
            prueba = self.carpeta / ".prueba_escritura"
            prueba.write_bytes(b"")
            prueba.unlink()
        except PermissionError as e:
            raise ErrorRegistro(
                f"No se puede escribir en el registro '{self.ruta}'. ¿Está abierto en Excel? "
                "Ciérralo y vuelve a ejecutar."
            ) from e
        except OSError as e:
            raise ErrorRegistro(f"No se puede acceder a la carpeta del registro '{self.carpeta}': {e}") from e

    def leer(self) -> list[FilaPrevia]:
        if not self.ruta.exists():
            return []
        try:
            wb = load_workbook(self.ruta, read_only=True)
        except Exception as e:
            raise ErrorRegistro(f"No se puede leer el registro '{self.ruta}': {e}") from e
        try:
            ws = wb[HOJA]
            filas = ws.iter_rows(values_only=True)
            cab = list(next(filas, []))
            if cab[: len(COLUMNAS)] != COLUMNAS:
                raise ErrorRegistro(f"El registro '{self.ruta}' no tiene las columnas esperadas.")
            i = {c: cab.index(c) for c in COLUMNAS}
            res = []
            for f in filas:
                if f is None or all(v is None for v in f):
                    continue
                res.append(
                    FilaPrevia(
                        fecha_proceso=str(f[i["Fecha proceso"]] or ""),
                        archivo=str(f[i["Archivo generado"]] or ""),
                        num_doc=str(f[i["Nº documento"]] or ""),
                        gln=str(f[i["GLN emisor"]] or ""),
                    )
                )
            return res
        except KeyError as e:
            raise ErrorRegistro(f"El registro '{self.ruta}' no tiene la hoja '{HOJA}'.") from e
        finally:
            wb.close()

    def preparar(self, filas: list[list]) -> Path:
        """Escribe el registro actualizado en un archivo temporal (aún no lo sustituye)."""
        if self.ruta.exists():
            wb = load_workbook(self.ruta)
            ws = wb[HOJA]
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = HOJA
            ws.append(COLUMNAS)
            for c, ancho in enumerate([19, 48, 18, 30, 15, 32, 14, 16, 8, 12, 40, 40], 1):
                ws.column_dimensions[get_column_letter(c)].width = ancho
            ws.freeze_panes = "A2"
        for fila in filas:
            ws.append(fila)
            for celda in ws[ws.max_row]:
                if isinstance(celda.value, str):
                    celda.number_format = "@"  # texto: que Excel no toque GLN ni Nº
        temporal = self.carpeta / f".{NOMBRE_ARCHIVO}.nuevo"
        wb.save(temporal)
        return temporal

    def confirmar(self, temporal: Path) -> None:
        try:
            os.replace(temporal, self.ruta)
        except OSError as e:
            temporal.unlink(missing_ok=True)
            raise ErrorRegistro(
                f"No se pudo guardar el registro '{self.ruta}' (¿se ha abierto en Excel?): {e}"
            ) from e
