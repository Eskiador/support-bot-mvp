"""Los ejemplos reales de ediwin: el resultado debe coincidir EXACTAMENTE."""

import re
from pathlib import Path

import pytest
from conftest import COMBINADO_1, COMBINADO_2, EJEMPLOS, contenido
from openpyxl import load_workbook
from pypdf import PdfReader

from divisor.analisis import analizar, leer_paginas
from divisor.proceso import procesar

P = "Confirmación Recepción_"

# (Nº doc, páginas del combinado, GLN emisor, nombre del archivo generado)
ESPERADO_LOTE1 = [
    ("28199", [1], "8480000009609", P + "MERCADONA RIBARROJA SECOS.pdf"),
    ("20265076540779", [2], "8430803111139", P + "7005232818.pdf"),
    ("1043946", [3], "8480011999999", P + "DESARROLLO DE MARCAS 1.pdf"),
    ("1041909", [4], "8480011999999", P + "DESARROLLO DE MARCAS 2.pdf"),
    ("20265076483596", [5, 6, 7], "8480015979997", P + "CARREFOUR 1.pdf"),
    ("20265076518284", [8, 9, 10], "8480015979997", P + "CARREFOUR 2.pdf"),
    ("201704004", [11, 12, 13], "8422410000005", P + "BON PREU 1.pdf"),
    ("201703984", [14], "8422410000005", P + "BON PREU 2.pdf"),
    ("201703973", [15], "8422410000005", P + "BON PREU 3.pdf"),
    ("201703835", [16, 17], "8422410000005", P + "BON PREU 4.pdf"),
    ("0690264428", [18, 19], "8480029069004", P + "PGC VALDEMORO.pdf"),
]

# Procesado DESPUÉS del lote 1 (Ribarroja ya existe sin número -> continúa en 2)
ESPERADO_LOTE2 = [
    ("15046", [1], "8480000010766", P + "MERCADONA VITORIA SECOS 1.pdf"),
    ("20310", [2], "8480000009265", P + "MERCADONA CIEMPOZUELOS SECOS 1.pdf"),
    ("22076", [3], "8480000009449", P + "MERCADONA ANTEQUERA SECOS 1.pdf"),
    ("10370", [4], "8480000010261", P + "MERCADONA GUADIX SECOS 1.pdf"),
    ("10362", [5], "8480000010261", P + "MERCADONA GUADIX SECOS 2.pdf"),
    ("14393", [6], "8480000011534", P + "MERCADONA 1153 1.pdf"),
    ("14683", [7], "8480000010766", P + "MERCADONA VITORIA SECOS 2.pdf"),
    ("21818", [8], "8480000009449", P + "MERCADONA ANTEQUERA SECOS 2.pdf"),
    ("19678", [9], "8480000009265", P + "MERCADONA CIEMPOZUELOS SECOS 2.pdf"),
    ("14209", [10], "8480000011534", P + "MERCADONA 1153 2.pdf"),
    ("19920", [11], "8480000009265", P + "MERCADONA CIEMPOZUELOS SECOS 3.pdf"),
    ("28193", [12], "8480000009609", P + "MERCADONA RIBARROJA SECOS 2.pdf"),
    ("12489", [13], "5606001007106", P + "IRMADONA 1.pdf"),
    ("12399", [14], "5606001007106", P + "IRMADONA 2.pdf"),
    ("12234", [15], "5606001007106", P + "IRMADONA 3.pdf"),
]

CABECERA = re.compile(r"^(Nº\.? [Cc]onfirmación: \S+ )?\d\d/\d\d/\d{4} \d\d:\d\d Página \d+$")


@pytest.mark.parametrize("combinado,esperado", [(COMBINADO_1, ESPERADO_LOTE1), (COMBINADO_2, ESPERADO_LOTE2)])
def test_division_del_combinado(cfg, combinado, esperado):
    a = analizar(combinado, cfg)
    assert [(d.num_doc, d.paginas, d.gln) for d in a.documentos] == [e[:3] for e in esperado]
    assert sum(len(d.paginas) for d in a.documentos) == a.total_paginas


def _sin_cabeceras(paginas):
    return [l for p in paginas for l in p if not CABECERA.match(l)]


@pytest.mark.parametrize("lote", ["lote1", "lote2"])
def test_cada_documento_coincide_con_su_descarga_individual(cfg, lote):
    """El contenido de cada documento del combinado = el PDF descargado uno a uno."""
    combinado = next((EJEMPLOS / lote).glob("report*.pdf"))
    a = analizar(combinado, cfg)
    individuales = {}
    for f in (EJEMPLOS / lote / "individuales").glob("*.pdf"):
        ai = analizar(f, cfg)  # combinado de un único documento
        assert len(ai.documentos) == 1
        individuales[ai.documentos[0].num_doc] = ai
    assert set(individuales) == {d.num_doc for d in a.documentos}
    for d in a.documentos:
        ind = individuales[d.num_doc]
        assert len(d.paginas) == ind.total_paginas
        assert _sin_cabeceras([a.textos[n - 1] for n in d.paginas]) == _sin_cabeceras(ind.textos)
        assert ind.documentos[0].gln == d.gln


def test_proceso_completo_de_los_dos_lotes(cfg, descargar):
    for combinado, esperado in [(COMBINADO_1, ESPERADO_LOTE1), (COMBINADO_2, ESPERADO_LOTE2)]:
        ruta = descargar(combinado)
        res = procesar(ruta, cfg)
        assert res.estado == "ok", res.mensaje
        assert [f["archivo"] for f in res.filas] == [e[3] for e in esperado]
        assert not ruta.exists()  # el combinado se ha movido a procesados

        lector = PdfReader(combinado)
        for num, paginas, _, nombre in esperado:
            salida = cfg.destino / nombre
            r = PdfReader(salida)
            assert len(r.pages) == len(paginas)
            # Páginas copiadas tal cual: mismo contenido gráfico byte a byte.
            for i, n in enumerate(paginas):
                assert r.pages[i].get_contents().get_data() == lector.pages[n - 1].get_contents().get_data()
            assert leer_paginas(salida) == [leer_paginas(combinado)[n - 1] for n in paginas]

    assert len(contenido(cfg.destino)) == len(ESPERADO_LOTE1) + len(ESPERADO_LOTE2)
    assert len(contenido(cfg.procesados)) == 2
    assert contenido(cfg.descargas) == []

    ws = load_workbook(cfg.registro / "registro_confirmaciones.xlsx").active
    filas = list(ws.iter_rows(values_only=True))[1:]
    assert [(f[2], f[1]) for f in filas] == [(e[0], e[3]) for e in ESPERADO_LOTE1 + ESPERADO_LOTE2]
    assert all(isinstance(f[4], str) and len(f[4]) == 13 for f in filas)  # GLN como texto


def test_simular_no_escribe_nada(cfg, descargar):
    ruta = descargar(COMBINADO_1)
    res = procesar(ruta, cfg, simular=True)
    assert res.estado == "simulado"
    assert [f["archivo"] for f in res.filas] == [e[3] for e in ESPERADO_LOTE1]
    assert contenido(cfg.destino) == []
    assert not cfg.registro.exists() and not cfg.procesados.exists()
    assert ruta.exists()


def test_combinado_de_un_unico_documento(cfg, descargar):
    """Un 'report' con un solo documento (descarga individual) también se procesa."""
    ind = EJEMPLOS / "lote1" / "individuales" / "report - 2026-10-05T120624.021.pdf"  # Carrefour, 3 págs
    res = procesar(descargar(ind), cfg)
    assert res.estado == "ok", res.mensaje
    assert contenido(cfg.destino) == [P + "CARREFOUR.pdf"]
    assert res.filas[0]["paginas"] == 3


def test_reprocesar_el_mismo_combinado_se_detiene(cfg, descargar):
    assert procesar(descargar(COMBINADO_1), cfg).estado == "ok"
    antes = contenido(cfg.destino)
    ruta = descargar(COMBINADO_1, "report - otra vez.pdf")
    res = procesar(ruta, cfg)
    assert res.estado == "error"
    assert "ya procesados" in res.mensaje and "28199" in res.mensaje
    assert contenido(cfg.destino) == antes
    assert ruta.exists()


# Fecha del documento y Nº de pedido (van al registro Excel), tal como los
# leía pdfplumber en la verificación inicial: el lector actual debe coincidir.
FECHA_Y_PEDIDO = {
    "28199": ("02/10/2026", "03078232"), "20265076540779": ("04/10/2026", "7005232818"),
    "1043946": ("04/10/2026", "67426"), "1041909": ("03/10/2026", "2087737"),
    "20265076483596": ("02/10/2026", "H336193227"), "20265076518284": ("02/10/2026", "H336193257"),
    "201704004": ("02/10/2026", "W210-832638"), "201703984": ("02/10/2026", "W210-832684"),
    "201703973": ("02/10/2026", "W210-832699"), "201703835": ("02/10/2026", "CC0-6313"),
    "0690264428": ("02/10/2026", "111683430000"),
    "15046": ("04/10/2026", "03151713"), "20310": ("04/10/2026", "03149885"),
    "22076": ("03/10/2026", "03132093"), "10370": ("03/10/2026", "03101332"),
    "10362": ("03/10/2026", "03131914"), "14393": ("02/10/2026", "03101341"),
    "14683": ("02/10/2026", "03101475"), "21818": ("02/10/2026", "03101338"),
    "19678": ("02/10/2026", "03078226"), "14209": ("02/10/2026", "03078233"),
    "19920": ("02/10/2026", "03101337"), "28193": ("02/10/2026", "03101340"),
    "12489": ("04/10/2026", "03149889"), "12399": ("04/10/2026", "03101333"),
    "12234": ("03/10/2026", "03132094"),
}


def test_fecha_y_pedido_de_cada_documento(cfg):
    leidos = {}
    for combinado in (COMBINADO_1, COMBINADO_2):
        for d in analizar(combinado, cfg).documentos:
            leidos[d.num_doc] = (d.fecha_documento, d.pedido)
    assert leidos == FECHA_Y_PEDIDO


# Lote 3 (05/10/2026 11:34): incluye CONSUM con el título "Aviso de expedición
# de mercancías" (variante de la plantilla A) y tres clientes nuevos.
COMBINADO_3 = EJEMPLOS / "lote3" / "report.pdf"
ESPERADO_LOTE3 = [
    ("24244", [1], "8480000009791", P + "MERCADONA SAN ISIDRO SECOS.pdf", "Confirmación Recepción Mercancías"),
    ("20205", [2], "8480000009265", P + "MERCADONA CIEMPOZUELOS SECOS 1.pdf", "Confirmación Recepción Mercancías"),
    ("15046", [3], "8480000010766", P + "MERCADONA VITORIA SECOS.pdf", "Confirmación Recepción Mercancías"),
    ("20310", [4], "8480000009265", P + "MERCADONA CIEMPOZUELOS SECOS 2.pdf", "Confirmación Recepción Mercancías"),
    ("14209", [5], "8480000011534", P + "MERCADONA 1153.pdf", "Confirmación Recepción Mercancías"),
    ("20265076330186", [6, 7, 8], "8480015979997", P + "CARREFOUR 1.pdf", "Aviso de recepción EAN"),
    ("20265076359278", [9, 10, 11], "8480015979997", P + "CARREFOUR 2.pdf", "Aviso de recepción EAN"),
    ("20265076292100", [12, 13, 14], "8480015979997", P + "CARREFOUR 3.pdf", "Aviso de recepción EAN"),
    ("5010124589", [15], "8413080000006", P + "TRANSGOURMET.pdf", "Aviso de recepción EAN"),
    ("1040590", [16], "8480011999999", P + "DESARROLLO DE MARCAS 1.pdf", "Aviso de recepción EAN"),
    ("1038687", [17], "8480011999999", P + "DESARROLLO DE MARCAS 2.pdf", "Aviso de recepción EAN"),
    ("201701577", [18], "8422410000005", P + "BON PREU 1.pdf", "Aviso de recepción EAN"),
    ("201701536", [19], "8422410000005", P + "BON PREU 2.pdf", "Aviso de recepción EAN"),
    ("0690261620", [20, 21], "8480029069004", P + "PGC VALDEMORO.pdf", "Aviso de recepción EAN"),
    ("00009200000357115", [22], "8414807000002", P + "CONSUM.pdf", "Aviso de expedición"),
]


def test_lote3_con_consum_y_clientes_nuevos(cfg, descargar):
    a = analizar(COMBINADO_3, cfg)
    assert [(d.num_doc, d.paginas, d.gln, d.tipo) for d in a.documentos] == [
        (e[0], e[1], e[2], e[4]) for e in ESPERADO_LOTE3
    ]
    ruta = descargar(COMBINADO_3)
    res = procesar(ruta, cfg)
    assert res.estado == "ok", res.mensaje
    assert [f["archivo"] for f in res.filas] == [e[3] for e in ESPERADO_LOTE3]
    lector = PdfReader(COMBINADO_3)
    for _, paginas, _, nombre, _ in ESPERADO_LOTE3:
        r = PdfReader(cfg.destino / nombre)
        assert [p.get_contents().get_data() for p in r.pages] == [
            lector.pages[n - 1].get_contents().get_data() for n in paginas
        ]


def test_lote3_tras_lote2_detecta_los_mercadona_repetidos(cfg, descargar):
    assert procesar(descargar(COMBINADO_2), cfg).estado == "ok"
    ruta = descargar(COMBINADO_3)
    res = procesar(ruta, cfg)
    assert res.estado == "error"
    for num in ("15046", "20310", "14209"):
        assert f"Nº {num}" in res.mensaje
    assert ruta.exists()


# Lote 4 (08/10/2026): Bon Preu manda DOS documentos con el mismo Nº 201704717
# (mismo albarán, distinto aviso de expedición; el segundo con cantidades
# negativas). Son documentos distintos y deben salir en PDFs distintos.
COMBINADO_4 = EJEMPLOS / "lote4" / "report.pdf"
ESPERADO_LOTE4 = [
    ("21029", [1], P + "MERCADONA CIEMPOZUELOS SECOS 1.pdf"),
    ("21345", [2], P + "MERCADONA CIEMPOZUELOS SECOS 2.pdf"),
    ("15362", [3], P + "MERCADONA 1153 1.pdf"),
    ("15359", [4], P + "MERCADONA 1153 2.pdf"),
    ("23059", [5], P + "MERCADONA ANTEQUERA SECOS 1.pdf"),
    ("23058", [6], P + "MERCADONA ANTEQUERA SECOS 2.pdf"),
    ("29954", [7], P + "MERCADONA RIBARROJA SECOS.pdf"),
    ("201705893", [8, 9], P + "BON PREU 1.pdf"),
    ("201705876", [10], P + "BON PREU 2.pdf"),
    ("201704717", [11, 12], P + "BON PREU 3.pdf"),
    ("201704717", [13], P + "BON PREU 4.pdf"),
    ("1045177", [14], P + "DESARROLLO DE MARCAS 1.pdf"),
    ("1044284", [15], P + "DESARROLLO DE MARCAS 2.pdf"),
    ("0690266519", [16, 17], P + "PGC VALDEMORO.pdf"),
    ("20265076687573", [18], P + "CARREFOUR.pdf"),
]


def test_lote4_mismo_numero_dos_documentos_distintos(cfg, descargar):
    a = analizar(COMBINADO_4, cfg)
    assert [(d.num_doc, d.paginas) for d in a.documentos] == [(e[0], e[1]) for e in ESPERADO_LOTE4]
    bon = [d for d in a.documentos if d.num_doc == "201704717"]
    assert len(bon) == 2 and bon[0].huella != bon[1].huella
    res = procesar(descargar(COMBINADO_4), cfg)
    assert res.estado == "ok", res.mensaje
    assert [f["archivo"] for f in res.filas] == [e[2] for e in ESPERADO_LOTE4]
    assert len(PdfReader(cfg.destino / (P + "BON PREU 3.pdf")).pages) == 2
    assert len(PdfReader(cfg.destino / (P + "BON PREU 4.pdf")).pages) == 1
    # Volver a procesar el mismo combinado: los dos Bon Preu están ya procesados
    ruta = descargar(COMBINADO_4, "report (1).pdf")
    res = procesar(ruta, cfg)
    assert res.estado == "error" and res.mensaje.count("Nº 201704717") == 2
    assert ruta.exists()


def test_la_huella_es_la_misma_en_cualquier_combinado(cfg):
    """El mismo documento descargado en combinados distintos (o suelto) da la
    misma huella: así se detecta que ya se procesó otro día."""
    def huellas(pdf):
        return {d.num_doc: d.huella for d in analizar(pdf, cfg).documentos}

    h2, h3 = huellas(COMBINADO_2), huellas(COMBINADO_3)
    for num in ("15046", "20310", "14209"):
        assert h2[num] == h3[num]
    h1 = huellas(COMBINADO_1)
    for f in (EJEMPLOS / "lote1" / "individuales").glob("*.pdf"):
        for num, h in huellas(f).items():
            assert h1[num] == h
